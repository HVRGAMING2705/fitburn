from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation(); prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
def add(i, title, lines):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame; tf.word_wrap = True
    for j, ln in enumerate(lines):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.text = ln; p.font.size = Pt(20); p.space_after = Pt(10)
add(0, "FitBurn — Calorie Burn Predictor", ["An AI/ML project by Vikranth Reddy", "B.Tech CSE (AI & ML), 2nd Year — Woxsen University", "Ridge regression that estimates workout calories in your browser"])
add(1, "The Problem", ["Fitness apps guess calorie burn with one-size-fits-all tables", "Two people doing the same workout can burn very different calories", "Age, weight, heart rate and intensity all change the number", "Users need a personal, explainable estimate — not a generic chart"])
add(2, "The Solution — FitBurn", ["Enter age, weight, duration, average heart rate and intensity (MET)", "A trained ridge regression model predicts calories instantly", "Runs fully in the browser — no backend, no waiting", "The same model weights power both the Python training and the web demo"])
add(3, "Tech Stack", ["Python + scikit-learn: data generation and model training", "Ridge regression (alpha = 1.0) on standardised features", "HTML, CSS and vanilla JavaScript for the interactive demo", "GitHub for source control, Vercel for free hosting"])
add(4, "Architecture", ["Synthetic dataset: 600 workouts, 5 features, seed 42", "Standardisation (mean / std) -> Ridge regression -> model.json", "model.json ships with the site; JavaScript applies the same formula", "Prediction is clipped to a realistic 30–1500 kcal range"])
add(5, "Live Demo", ["Try it: enter your workout details and press Estimate", "Example: 30 yrs, 75 kg, 90 min, 155 bpm, 8.5 MET -> about 964 kcal", "Change any input and the estimate updates instantly", "Demo: https://fitburn.vercel.app"])
add(6, "Results", ["5-fold cross-validation R-squared: 0.907", "Mean absolute error: 89.8 kcal (RMSE 116.3 kcal)", "Baseline (always guessing the average): MAE 310.6 kcal", "The model is about 3.5x better than the baseline guess", "Browser demo matches Python predictions to within 0.01 kcal"])
add(7, "Future Scope", ["Train on real wearable data (Fitbit, Apple Watch exports)", "Add workout type and user history for personalisation", "Compare gradient boosting and small neural networks", "Add a daily calorie-goal tracker and progress charts"])
add(8, "Thank You", ["FitBurn — Calorie Burn Predictor", "Vikranth Reddy, B.Tech CSE (AI & ML), Woxsen University", "Source: github.com/HVRGAMING2705/fitburn", "Questions welcome!"])
prs.save("/home/hatch/workspace/goals/daily-ai-ml-project-factory/files/AIML_FitBurn_2026-10-30.pptx")
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
